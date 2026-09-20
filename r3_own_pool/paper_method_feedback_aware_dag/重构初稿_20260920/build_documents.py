"""Build review DOCX with native Word equations from the Chinese Markdown draft."""
from pathlib import Path
import subprocess,re,json,hashlib,datetime
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).parent
D=Document();sec=D.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
for name in ['Normal','Body Text','First Paragraph','Title','Heading 1','Heading 2','Heading 3','Heading 4','Caption']:
 if name not in D.styles:continue
 s=D.styles[name];s.font.name='Liberation Serif';s.font.size=Pt(10.5);s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC');s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.2
 if name.startswith('Heading'):
  level=int(name[-1]);s.font.size=Pt({1:18,2:15,3:12,4:11}[level]);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(14)
D.save(P/'排版模板.docx')
subprocess.run(['pandoc',str(P/'论文完整初稿.md'),'-f','markdown+tex_math_dollars','-t','docx','--reference-doc='+str(P/'排版模板.docx'),'-o',str(P/'论文完整初稿.docx')],check=True)
d=Document(P/'论文完整初稿.docx')
for p in d.paragraphs:
 if p.text.startswith('表 '):p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
# Rebuild tables using standard Word elements; retain every cell's text.
for old in list(d.tables):
 rows=[[c.text for c in row.cells] for row in old.rows]
 new=d.add_table(rows=len(rows),cols=len(rows[0]))
 new.style='Table Grid'
 for ri,row in enumerate(rows):
  for ci,value in enumerate(row):new.cell(ri,ci).text=value
 old._tbl.addprevious(new._tbl)
 old._tbl.getparent().remove(old._tbl)
for t in d.tables:
 t.autofit=False
 tw=t._tbl.tblPr.find(qn("w:tblW"));tw.set(qn("w:type"),"dxa");tw.set(qn("w:w"),str(Cm(16.8).twips))
 n=len(t.rows[0].cells)
 if not len(t.columns):
  for _ in range(n):t._tbl.tblGrid.append(OxmlElement("w:gridCol"))
 widths=[16.8/n]*n
 if n==8:widths=[3.0,1.35,1.7,1.85,2.1,2.0,1.6,1.7]
 elif n==7:widths=[3.0,1.45,1.65,1.65,1.7,3.0,2.85]
 elif n==6:widths=[2.9,1.5,2.6,2.7,3.5,3.1]
 elif n==5:widths=[3.8,2.2,2.7,3.5,4.6]
 elif n==4:widths=[4.0,3.0,4.7,5.1]
 elif n==3:widths=[4.2,5.0,7.6]
 total=sum(widths);widths=[x*16.8/total for x in widths]
 for i,col in enumerate(t.columns):col.width=Cm(widths[i])
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for i,cell in enumerate(row.cells):
   cell.width=Cm(widths[i]);cell.vertical_alignment=1
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');cell._tc.get_or_add_tcPr().append(sh)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=1.0
    for run in p.runs:
     run.font.size=Pt(8 if n>=7 else 9);run.font.bold=ri==0
     run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
for sec in d.sections:
 p=sec.footer.paragraphs[0];p.alignment=1;r=p.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
d.core_properties.title='面向复杂任务的反馈感知多模型自适应有向无环图路由';d.core_properties.subject='2026-09-20 证据重构初稿';d.save(P/'论文完整初稿.docx')
root=P.parents[1]
refs=set(re.findall(r'`(static_dag_v0/[^`]+\.(?:json|jsonl|npz|py))`',(P/'PAPER_EVIDENCE_MAP.md').read_text()))
refs.update(str(x.relative_to(root))for x in P.parent.glob('*.md'))
refs.update(['static_dag_v0/dynamic_v2_policy.py','static_dag_v0/dynamic_v2_test.py','static_dag_v0/live_static_dynamic.py','static_dag_v0/budget_surface.py','static_dag_v0/controlled_fault_analyze.py','static_dag_v0/scale_up_analyze.py'])
manifest=[]
for f in sorted(refs):
 p=root/f
 if not p.exists():raise FileNotFoundError(f)
 kind='历史论文/方法说明' if 'paper_method' in f else ('统计或执行实现'if p.suffix=='.py'else '实验数据/结果/协议')
 manifest.append({'path':f,'category':kind,'bytes':p.stat().st_size,'mtime_utc':datetime.datetime.fromtimestamp(p.stat().st_mtime,datetime.timezone.utc).isoformat(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(P/'MATERIAL_INVENTORY.json').write_text(json.dumps({'scope':'Core evidence used in this reconstruction, not an assertion of line-by-line review of every historical log.','files':manifest},ensure_ascii=False,indent=2))
text=(P/'论文完整初稿.md').read_text();heads=[l for l in text.splitlines()if l.startswith('#')];bad=[l for l in heads if re.search('[A-Za-z]',l)]
assert not bad,bad
assert len(d.tables)==18
assert sum(p.text.startswith('待我确认项') for p in d.paragraphs)==8
from zipfile import ZipFile
with ZipFile(P/'论文完整初稿.docx')as z:xml=z.read('word/document.xml').decode()
checks={'chinese_characters':len(re.findall('[\u4e00-\u9fff]',text)),'tables':len(d.tables),'headings':len(heads),'english_headings':bad,'native_equations':xml.count('<m:oMath>'),'chapter_review_notes':8,'evidence_inventory_files':len(manifest)}
(P/'DOCUMENT_CHECKS.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2));print(checks)
