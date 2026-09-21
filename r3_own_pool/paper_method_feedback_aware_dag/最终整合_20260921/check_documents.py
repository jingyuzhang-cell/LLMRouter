from pathlib import Path
import json,re,xml.etree.ElementTree as ET
from zipfile import ZipFile
from docx import Document
P=Path(__file__).parent;R=P.parents[1];text=(P/'论文完整草稿.md').read_text()
heads=[l for l in text.splitlines() if l.startswith('#')]
assert not [l for l in heads if re.search('[A-Za-z]',l)]
assert len(re.findall(r'^### 3\.\d+ ',text,re.M))==10
assert len(re.findall(r'^### 5\.\d+ ',text,re.M))==8
assert not any(ord(c)<32 and c not in '\n\t' for c in text)
r=json.loads((P.parent/'重构初稿_20260920/DERIVED_EVIDENCE.json').read_text())['routing']['Q']
assert abs((r['FrozenNodeRouter']-r['AlwaysLarge'])/(r['NodeOracle']-r['AlwaysLarge'])-33/68)<1e-10
x=json.loads((R/'static_dag_v0/exact_optimality_audit/AUDIT.json').read_text());assert x['Q_exact_oracle']==.305 and x['Q_always_large']==.18 and x['Q_best_fixed']==.205
x=json.loads((R/'static_dag_v0/confirmation_500/CONF500_RESULTS.json').read_text());assert x['help']==1 and x['harm']==4 and len(x['switches'])==22 and x['GAP_recovery']==-.1111
with ZipFile(P/'论文完整草稿.docx') as z: xml=z.read('word/document.xml').decode()
d=Document(P/'论文完整草稿.docx');assert len(d.tables)==15,len(d.tables)
pages=ET.parse('/tmp/final_paper_bbox.html').getroot().findall('.//{http://www.w3.org/1999/xhtml}page');outside=[]
for n,page in enumerate(pages,1):
 for w in page.findall('.//{http://www.w3.org/1999/xhtml}word'):
  if float(w.attrib['xMin'])<0 or float(w.attrib['xMax'])>float(page.attrib['width'])+1 or float(w.attrib['yMax'])>float(page.attrib['height'])+1:outside.append((n,w.text))
assert not outside,outside
checks={'chinese_characters':len(re.findall('[\u4e00-\u9fff]',text)),'all_headings_chinese':True,'method_subsections':10,'experiment_research_questions':8,'docx_tables':len(d.tables),'native_math_elements':xml.count('<m:oMath>'),'pdf_pages':len(pages),'pdf_out_of_page_words':outside,'numeric_spot_checks_passed':True,'pending_source_items':[],'resolved_source_items':{'propagated reasoning headroom 4.1pp':'static_dag_v0/propagated_row_oracle_audit/REASONER_ROW_ORACLE_AUDIT.json (5.00/4.50/5.40/5.06pp; 4.1pp void)','complete paired D0-D4 diagnostic scores':'static_dag_v0/structure_aware_experiment/EVIDENCE_DIAGNOSIS_AUDIT.json (B-arm Q=0/25; coverage claim non-reproducible; D0/D1/C_Gold no artifacts; ERv2 8/8/5/8%)'},'scope':'Document and frozen-source consistency checks; no new experiment or model calls'}
(P/'CONSISTENCY_CHECKS.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2));print(json.dumps(checks,ensure_ascii=False))
